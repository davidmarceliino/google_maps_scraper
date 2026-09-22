import csv

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

COLUMNS = ["Nome", "Telefone", "Endereço", "Cidade", "Estado", "CEP", "Site", "Google Maps"]

HEADER_TO_FIELD = {
    "Nome": "name",
    "Telefone": "phone",
    "Endereço": "address",
    "Cidade": "city",
    "Estado": "state",
    "CEP": "cep",
    "Site": "website",
    "Google Maps": "google_maps",
}


def _extract(business):
    return [(business or {}).get(HEADER_TO_FIELD[h], "") or "" for h in COLUMNS]


def export_excel(businesses, path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Empresas"

    ws.append(COLUMNS)
    fill = PatternFill("solid", fgColor="1F4E79")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center")

    for biz in businesses:
        ws.append(_extract(biz))

    widths = [40, 18, 50, 20, 8, 12, 35, 45]
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[chr(64 + i)].width = width

    ws.freeze_panes = "A2"
    wb.save(path)


def export_csv(businesses, path):
    with open(path, "w", newline="", encoding="utf-8-sig") as fh:
        writer = csv.writer(fh, delimiter=";")
        writer.writerow(COLUMNS)
        for biz in businesses:
            writer.writerow(_extract(biz))